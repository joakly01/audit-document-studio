"""Create original-format files for the explicitly constructed initial catalog only."""
import sys,json,csv,io,hashlib
from pathlib import Path
from xml.sax.saxutils import escape
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.pagesizes import A4
from docx import Document
from docx.shared import Pt
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment
base=Path(__file__).resolve().parents[1];out=base/'dist/source-files';out.mkdir(parents=True,exist_ok=True)
docs=json.loads(Path(sys.argv[1]).read_text());manifest={}
pdfmetrics.registerFont(TTFont('AuditChinese','/usr/share/fonts/truetype/wqy/wqy-zenhei.ttc',subfontIndex=0))
for d in docs:
 ext=d['name'].rsplit('.',1)[-1].lower();path=out/(d['id']+'.'+ext)
 text=d.get('text') or (d['name'].rsplit('.',1)[0]+'\n资料归属：'+d.get('unit','')+'。\n事项：'+d.get('topic','资料整理')+'。\n本资料用于记录审批、合同履约与资金支付情况。请核对相关合同、凭证和验收记录，缺少材料应继续补充核实。')
 title=d.get('title') or d['name'];meta=d.get('unit','')+'　'+d.get('date','')
 if ext=='pdf':
  normal=ParagraphStyle('body',fontName='AuditChinese',fontSize=11,leading=20,wordWrap='CJK',spaceAfter=10)
  heading=ParagraphStyle('heading',parent=normal,fontSize=18,leading=28,spaceAfter=16)
  content=[Paragraph(escape(title),heading),Paragraph(escape(meta),normal),Spacer(1,12)]+[Paragraph(escape(line) or ' ',normal) for line in text.splitlines()]
  SimpleDocTemplate(str(path),pagesize=A4,rightMargin=54,leftMargin=54,topMargin=54,bottomMargin=54,title=title,author=d.get('unit','')).build(content)
 elif ext=='docx':
  doc=Document();doc.styles['Normal'].font.name='微软雅黑';doc.styles['Normal'].font.size=Pt(11);doc.add_heading(title,0);doc.add_paragraph(meta)
  for line in text.splitlines():doc.add_paragraph(line)
  doc.save(path)
 elif ext in ['xlsx','xls']:
  if ext=='xls':continue
  wb=Workbook();ws=wb.active;ws.title='资料明细';matrix=[d.get('columns') or ['项目','内容'],*(d.get('rows') or [[title,text]])]
  for row in matrix:ws.append(row)
  for cell in ws[1]:cell.font=Font(name='微软雅黑',bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='42678F')
  ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
  for column in ws.columns:
   ws.column_dimensions[column[0].column_letter].width=28
   for cell in column:cell.alignment=Alignment(vertical='top',wrap_text=True)
  wb.save(path)
 elif ext=='csv':
  with path.open('w',encoding='utf-8-sig',newline='') as f:csv.writer(f).writerows([d.get('columns') or ['事项','内容'],*(d.get('rows') or [[title,text]])])
 elif ext in ['txt','md','json']:path.write_text(text)
 else:continue
 mime={'pdf':'application/pdf','docx':'application/vnd.openxmlformats-officedocument.wordprocessingml.document','xlsx':'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet','csv':'text/csv','txt':'text/plain','md':'text/markdown','json':'application/json'}[ext]
 manifest[d['id']]={'path':'source-files/'+path.name,'name':d['name'],'mime':mime,'size':path.stat().st_size,'sha256':hashlib.sha256(path.read_bytes()).hexdigest()}
(base/'dist/source-assets.js').write_text("'use strict';\n// Original-format assets for the constructed initial catalog, never substitutes for uploaded files.\nconst AUDIT_SOURCE_ASSETS=Object.freeze("+json.dumps(manifest,ensure_ascii=False,separators=(',',':'))+");\nfor(const d of state.docs){const asset=AUDIT_SOURCE_ASSETS[d.id];if(asset&&!d.byteSize){d.byteSize=asset.size;d.size=libFormat(asset.size);}}\n")
print('Built',len(manifest),'source files; bytes',sum(x['size'] for x in manifest.values()))
