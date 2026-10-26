"""Generate consistent HTML, PDF and editable DOCX from the reviewed Markdown."""
from pathlib import Path
import sys, re, base64, json, hashlib, html as esc, unicodedata
sys.path.insert(0,'/tmp/audit-plan-doc-libs')
import markdown
from lxml import html, etree
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from playwright.sync_api import sync_playwright
from pypdf import PdfReader

ROOT=Path(__file__).resolve().parent
NAME='文档挖掘助手-总体优化设计方案-V1.0'
FILES=sorted(ROOT.glob('[0-9][0-9]-*.md'))
ASSETS=ROOT/'evidence';ASSETS.mkdir(exist_ok=True)
SVG='''<svg xmlns="http://www.w3.org/2000/svg" width="1200" height="330" viewBox="0 0 1200 330">
<defs><marker id="a" markerWidth="8" markerHeight="8" refX="7" refY="4" orient="auto"><path d="M0 0L8 4L0 8" fill="none" stroke="#6b829b"/></marker></defs>
<rect width="1200" height="330" fill="#f4f7fb" rx="16"/>
<g font-family="Noto Sans CJK SC,sans-serif" text-anchor="middle">
<text x="600" y="32" font-size="18" font-weight="600" fill="#20374c">审计人员的日常工作路径</text>
'''
steps=[('添加资料','明确本次范围'),('设置分析内容','字段／提示词／规则'),('开始分析','真实回执与进度'),('查看结果','本工具专用成果'),('核验依据','原文与人工意见'),('导出成果','明确范围与来源')]
for i,(title,sub) in enumerate(steps):
    x=20+i*198
    SVG+=f'<rect x="{x}" y="60" width="170" height="92" rx="10" fill="white" stroke="#cbd8e8"/><text x="{x+85}" y="94" font-size="21" fill="#174f96" font-weight="600">{title}</text><text x="{x+85}" y="124" font-size="15" fill="#52657b">{sub}</text>'
    if i<5: SVG+=f'<path d="M{x+172} 106H{x+194}" stroke="#6b829b" fill="none" marker-end="url(#a)"/>'
SVG+='''<path d="M700 156V180H305V156" stroke="#6b829b" fill="none" marker-end="url(#a)"/><rect x="394" y="164" width="233" height="30" rx="15" fill="#f4f7fb"/><text x="510" y="185" font-size="15" fill="#174f96">再次分析：先编辑，再新建</text>
<rect x="25" y="220" width="360" height="77" rx="9" fill="#e8eff9"/><text x="205" y="249" font-size="18" fill="#20374c" font-weight="600">资料管理与模板／提示词</text><text x="205" y="276" font-size="15" fill="#52657b">管理资料归属 · 复用业务经验</text>
<rect x="420" y="220" width="360" height="77" rx="9" fill="#e8eff9"/><text x="600" y="249" font-size="18" fill="#20374c" font-weight="600">工具内分析记录</text><text x="600" y="276" font-size="15" fill="#52657b">离开后找回进度与成果</text>
<rect x="815" y="220" width="360" height="77" rx="9" fill="#e8eff9"/><text x="995" y="249" font-size="18" fill="#20374c" font-weight="600">任务中心</text><text x="995" y="276" font-size="15" fill="#52657b">跨工具查成果 · 集中看运行</text></g></svg>'''
(ASSETS/'业务主流程.svg').write_text(SVG)

def data_url(p):
    mime='image/svg+xml' if p.suffix=='.svg' else 'image/png'
    return 'data:'+mime+';base64,'+base64.b64encode(p.read_bytes()).decode()

chapters=[]
def column_weights(heads):
    n=len(heads)
    if n==2:return [30,70]
    if n==3:return [23,42,35]
    if n==4:return [16,25,30,29]
    if n==5:
        return [9,28,15,28,20] if '发现与证据' in heads[1] else [12,17,25,25,21]
    return [100/n]*n

for file in FILES:
    key='c'+file.name[:2]
    raw=file.read_text()
    source=markdown.markdown(raw,extensions=['tables','fenced_code','sane_lists'])
    root=html.fragment_fromstring(source,create_parent='section')
    title=root.find('h1').text_content()
    root.set('id',key);root.set('class','chapter')
    for j,h in enumerate(root.xpath('.//h1|.//h2|.//h3')):
        h.set('id',f'{key}-h{j}')
    for table in root.xpath('.//table'):
        heads=[n.text_content() for n in table.xpath('./thead/tr/th')]
        group=etree.Element('colgroup')
        for w in column_weights(heads):
            col=etree.SubElement(group,'col');col.set('style',f'width:{w}%')
        table.insert(0,group)
    for code in root.xpath('.//pre/code[@class="language-mermaid"]'):
        figure=html.fromstring('<figure><img alt="日常分析与结果核验流程"/><figcaption>日常分析、历史返回与再次分析的关系</figcaption></figure>')
        figure.find('img').set('src','evidence/业务主流程.svg')
        code.getparent().getparent().replace(code.getparent(),figure)
    chapters.append({'key':key,'file':file,'title':title,'root':root})

css='''
:root{--ink:#21364c;--muted:#52657b;--blue:#235da8;--line:#d8e2ee;--paper:#fff;--page:#f1f4f8}
*{box-sizing:border-box}[hidden]{display:none!important}body{margin:0;color:var(--ink);background:var(--page);font:16px/1.8 "Noto Sans CJK SC","Microsoft YaHei",sans-serif}
a{color:#1b58a4;text-decoration:none;overflow-wrap:anywhere}a:hover{text-decoration:underline}button,input{font:inherit}
.nav{position:fixed;inset:0 auto 0 0;width:264px;background:#182e43;color:#d6e2f0;padding:28px 20px;overflow:auto}.nav strong{display:block;color:white;font-size:19px;margin-bottom:6px}.nav small{display:block;color:#bdcee1;margin-bottom:24px}.nav a{display:block;color:#e1eaf3;font-size:14px;padding:9px 10px;border-radius:6px;line-height:1.55}.nav a:hover{background:#2a4663;text-decoration:none}.nav input{width:100%;padding:8px;margin-bottom:16px;border-radius:6px;border:1px solid #52657b;background:#fff;color:var(--ink);font-size:14px}
.toolbar{position:sticky;top:0;background:#ffffffed;border-bottom:1px solid var(--line);margin-left:264px;padding:10px 28px;display:flex;justify-content:space-between;z-index:5}.toolbar button{background:white;color:var(--blue);padding:5px 18px;border:1px solid #b9cee7;border-radius:6px;cursor:pointer}.toolbar span{font-size:14px;color:var(--muted)}
main{margin:24px 30px 64px 294px;max-width:1180px}.cover,.chapter{background:white;border:1px solid var(--line);padding:46px 50px;margin-bottom:24px;border-radius:10px;scroll-margin-top:70px}.cover{padding:58px 50px 44px;border-top:8px solid var(--blue)}.eyebrow{font-size:14px;letter-spacing:2px;color:var(--blue);font-weight:600}.cover h1{font-size:40px;line-height:1.5;margin:18px 0}.cover .subtitle{font-size:20px;color:var(--muted)}.cover .lead{padding:18px 0;font-size:17px;max-width:760px}.cover .cover-meta{border-top:1px solid var(--line);padding-top:22px;font-size:14px;color:var(--muted)}
h1{font-size:28px;line-height:1.5;padding-bottom:16px;border-bottom:2px solid var(--blue);margin:0 0 22px}h2{font-size:21px;line-height:1.5;margin:30px 0 13px;color:#194f88}h3{font-size:18px;margin:24px 0 10px}h1,h2,h3{scroll-margin-top:76px}p{margin:10px 0 14px}ul,ol{padding-left:1.6em}li{padding-bottom:4px}strong{font-weight:650}blockquote{margin:15px 0;padding:10px 18px;border-left:4px solid #4078bb;background:#f0f5fc;color:#465d75}blockquote p{margin:2px 0}
table{width:100%;border-collapse:collapse;table-layout:fixed;margin:18px 0 23px;font-size:14px;line-height:1.7}th,td{border:1px solid var(--line);padding:11px 12px;vertical-align:top;overflow-wrap:anywhere}th{background:#eaf1fa;color:#194f88;font-weight:600;text-align:left}tr:nth-child(even) td{background:#f9fbfe}code{font:0.91em/1.6 Consolas,"Noto Sans Mono CJK SC",monospace;color:#23516e;background:#f0f4f9;padding:1px 3px;border-radius:3px;overflow-wrap:anywhere}pre{white-space:pre-wrap;padding:14px;background:#f1f5fa;font-size:13px;overflow-wrap:anywhere}figure{margin:22px 0;text-align:center}figure img,p>img{display:block;max-width:100%;height:auto;border:1px solid var(--line);border-radius:6px}figcaption{font-size:13px;color:#52657b;margin-top:8px}img{max-width:100%}.chapter .chapter-number{font-size:12px;letter-spacing:2px;color:#527498}
@media(max-width:1100px){.nav{width:218px;padding:20px 12px}.toolbar{margin-left:218px}main{margin:16px 16px 40px 234px}.cover,.chapter{padding:30px 24px}table{font-size:13px}th,td{padding:8px}}
@media(max-width:740px){.nav{position:static;width:auto;max-height:240px}.nav strong,.nav small{display:inline-block;margin-right:10px}.nav input{margin:8px 0}.nav a{display:inline-block;font-size:13px;padding:5px}.toolbar{margin-left:0;padding:8px 14px}main{margin:12px}.cover,.chapter{padding:26px 18px}.cover h1{font-size:30px}h1{font-size:24px}h2{font-size:20px}table{display:block;overflow-x:auto;width:100%}th,td{min-width:115px}}
@page{size:A4;margin:16mm 15mm 17mm}
@media print{body{background:white;font-size:10.5pt;line-height:1.65;color:#21364c;-webkit-print-color-adjust:exact;print-color-adjust:exact}.nav,.toolbar{display:none}main{margin:0;max-width:none}.cover,.chapter{border:0;border-radius:0;margin:0;padding:0}.cover{height:250mm;page-break-after:always;border-top:4mm solid #235da8;padding-top:26mm}.cover h1{font-size:32pt;border:0;margin:16mm 0 10mm;line-height:1.55}.cover .subtitle{font-size:17pt}.cover .lead{font-size:13pt;margin-top:14mm}.cover .cover-meta{margin-top:24mm;font-size:10pt}.cover img{margin-top:12mm}.chapter{page-break-before:always}h1{font-size:20pt;margin:0 0 16pt;padding-bottom:10pt}h2{font-size:14pt;margin:18pt 0 9pt}h3{font-size:12pt;margin:13pt 0 7pt}h1,h2,h3{break-after:avoid}p{margin:7pt 0 9pt;orphans:3;widows:3}li{padding-bottom:2pt}blockquote{padding:8pt 10pt;font-size:9.5pt}table{font-size:8.5pt;line-height:1.55;margin:10pt 0 14pt;display:table}th,td{padding:6pt;min-width:0}thead{display:table-header-group}tr{break-inside:avoid}code{font-size:.92em}figure{break-inside:avoid;margin:12pt 0}figure img,p>img{max-height:160mm;object-fit:contain;margin:auto}p:has(>img){break-inside:avoid}a{color:#235da8}pre{font-size:8.5pt}ul,ol{margin:6pt 0 10pt}figcaption{font-size:8.5pt}}
'''
cover='''<section class="cover"><div class="eyebrow">审计业务 · 产品设计 · 研发实施</div><h1>文档挖掘助手<br>总体优化设计方案</h1><p class="subtitle">从资料处理到成果核验的完整工作路径</p><p class="lead">统一操作习惯，保留业务差异。<br>以现有前端源码为依据，明确已实现、待接入和需补充的能力。</p><img src="evidence/业务主流程.svg" alt="完整工作流程"/><div class="cover-meta">V1.0 · 2026年10月25日<br>评审基线：20261024.80 · 11项分析工具及公共模块<br>交付范围：优化设计、UI规范、后端对接、审查与验收</div></section>'''
nav=''.join(f'<a href="#{c["key"]}">{c["file"].name[:2]}　{esc.escape(c["title"].replace("文档挖掘助手 · ",""))}</a>' for c in chapters)
allroots=[]
for c in chapters:
    clone=html.fromstring(etree.tostring(c['root'],encoding='unicode',method='html'))
    for a in clone.xpath('.//a[@href]'):
        href=a.get('href')
        matches=[x for x in chapters if href==x['file'].name]
        if matches:a.set('href','#'+matches[0]['key'])
    allroots.append(etree.tostring(clone,encoding='unicode',method='html'))
page='<!doctype html><html lang="zh-CN"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>'+NAME+'</title><style>'+css+'</style></head><body><nav class="nav" aria-label="方案章节"><strong>文档挖掘助手</strong><small>总体优化设计 · V1.0</small><input aria-label="查找章节" placeholder="查找章节" oninput="document.querySelectorAll(\'.nav a\').forEach(a=>a.hidden=!a.textContent.includes(this.value))">'+nav+'</nav><div class="toolbar"><span>设计方案与开发交接 · 2026-10-25</span><button onclick="window.print()">打印 / 保存 PDF</button></div><main>'+cover+''.join(allroots)+'</main></body></html>'
dom=html.fromstring(page)
for img in dom.xpath('//img[@src]'):
    p=ROOT/img.get('src')
    if p.exists():img.set('src',data_url(p))
htmlpath=ROOT/(NAME+'.html')
htmlpath.write_text('<!doctype html>\n'+etree.tostring(dom,encoding='unicode',method='html'))

with sync_playwright() as p:
    b=p.chromium.launch(args=['--no-sandbox']);g=b.new_page(viewport={'width':1200,'height':330},device_scale_factor=1.5)
    g.goto((ASSETS/'业务主流程.svg').as_uri());g.screenshot(path=str(ASSETS/'业务主流程.png'))
    g.set_viewport_size({'width':1440,'height':1000});g.goto(htmlpath.as_uri());g.evaluate('document.fonts.ready');g.screenshot(path=str(ASSETS/'方案阅读版预览.png'))
    g.pdf(path=str(ROOT/(NAME+'.pdf')),format='A4',print_background=True,prefer_css_page_size=True,display_header_footer=True,
      header_template='<div></div>',footer_template='<div style="font-size:8px;color:#64748b;width:100%;margin:0 15mm;display:flex;justify-content:space-between"><span>文档挖掘助手 · 总体优化设计方案 V1.0</span><span><span class="pageNumber"></span> / <span class="totalPages"></span></span></div>',outline=True,tagged=True)
    b.close()

# Editable Word edition from the same parsed content.
doc=Document();sec=doc.sections[0];sec.page_width=Cm(21);sec.page_height=Cm(29.7)
sec.top_margin=Cm(1.8);sec.bottom_margin=Cm(1.8);sec.left_margin=Cm(1.6);sec.right_margin=Cm(1.6)
normal=doc.styles['Normal'];normal.font.name='Microsoft YaHei';normal.font.size=Pt(10.5);normal.font.color.rgb=RGBColor.from_string('21364C')
normal.element.rPr.rFonts.set(qn('w:eastAsia'),'微软雅黑');normal.paragraph_format.line_spacing=1.5;normal.paragraph_format.space_after=Pt(7)
for i,sz in [(1,20),(2,14),(3,12)]:
    st=doc.styles['Heading '+str(i)];st.font.name='Microsoft YaHei';st.font.size=Pt(sz);st.font.color.rgb=RGBColor.from_string('194F88');st.element.rPr.rFonts.set(qn('w:eastAsia'),'微软雅黑');st.paragraph_format.keep_with_next=True
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('文档挖掘助手 · V1.0    ').font.size=Pt(8)
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
doc.core_properties.title=NAME;doc.core_properties.subject='业务流程、UI规范、原前端对照与后端开发清单';doc.core_properties.author='文档挖掘助手项目'
doc.add_paragraph('审计业务 · 产品设计 · 研发实施').runs[0].font.color.rgb=RGBColor.from_string('235DA8')
p=doc.add_paragraph();p.paragraph_format.space_before=Pt(60)
r=p.add_run('文档挖掘助手\n总体优化设计方案');r.bold=True;r.font.size=Pt(30)
doc.add_paragraph('从资料处理到成果核验的完整工作路径').runs[0].font.size=Pt(16)
doc.add_paragraph('统一操作习惯，保留业务差异。\n以原前端源码为依据，明确已实现、待接入和需补充的能力。')
doc.add_picture(str(ASSETS/'业务主流程.png'),width=Cm(17.6))
doc.add_paragraph('V1.0｜2026-10-25\n评审基线：20261024.80\n11项分析工具及公共模块\n优化设计 · UI规范 · 后端对接 · 审查验收')

def inline(p,node,bold=False,italic=False):
    def run(text,b=False,i=False):
        if not text:return
        r=p.add_run(text);r.bold=b;r.italic=i
    if node.text:run(node.text,bold,italic)
    for e in node:
        if e.tag=='a':
            link=OxmlElement('w:hyperlink');href=e.get('href','');
            rel=p.part.relate_to(href,RT.HYPERLINK,is_external=True);link.set(qn('r:id'),rel)
            r=OxmlElement('w:r');prop=OxmlElement('w:rPr');col=OxmlElement('w:color');col.set(qn('w:val'),'235DA8');prop.append(col);r.append(prop)
            t=OxmlElement('w:t');t.text=e.text_content();r.append(t);link.append(r);p._p.append(link)
        elif e.tag=='br':p.add_run().add_break()
        elif e.tag=='img':
            path=ROOT/e.get('src','');path=ASSETS/'业务主流程.png' if path.suffix=='.svg' else path
            if path.exists():p.add_run().add_picture(str(path),width=Cm(17.5))
        else:inline(p,e,bold or e.tag in ['strong','b'],italic or e.tag in ['em','i'])
        if e.tail:run(e.tail,bold,italic)

def block(e,target=doc):
    tag=e.tag
    if tag in ['h1','h2','h3']:
        p=target.add_paragraph(style='Heading '+tag[1]);inline(p,e);return
    if tag=='p':
        p=target.add_paragraph();inline(p,e);return
    if tag in ['ul','ol']:
        for i,li in enumerate(e.findall('li')):
            p=target.add_paragraph(style='List Bullet' if tag=='ul' else 'List Number');inline(p,li)
        return
    if tag=='blockquote':
        for pe in e:
            p=target.add_paragraph();p.paragraph_format.left_indent=Cm(.4);inline(p,pe)
            for r in p.runs:r.font.size=Pt(9.5);r.font.color.rgb=RGBColor.from_string('52657B')
        return
    if tag=='table':
        rows=e.xpath('./thead/tr|./tbody/tr|./tr');n=max(len(row) for row in rows)
        t=target.add_table(rows=0,cols=n);t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False;t.style='Table Grid'
        # Narrow identifiers; leave descriptions room.
        heads=[x.text_content() for x in rows[0]]
        weights=column_weights(heads)
        total=sum(weights);widths=[17.7*w/total for w in weights]
        for j,w in enumerate(widths):t.columns[j].width=Cm(w)
        for i,row in enumerate(rows):
            cells=t.add_row().cells
            pr=t.rows[-1]._tr.get_or_add_trPr();no=OxmlElement('w:cantSplit');pr.append(no)
            if i==0:repeat=OxmlElement('w:tblHeader');pr.append(repeat)
            for j,el in enumerate(row):
                c=cells[j];c.width=Cm(widths[j]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.TOP
                p=c.paragraphs[0];p.paragraph_format.space_after=Pt(3);p.paragraph_format.space_before=Pt(3);p.paragraph_format.line_spacing=1.25;inline(p,el,bold=i==0)
                for r in p.runs:r.font.size=Pt(8.5)
                if i==0 or i%2==0:
                    sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'EAF1FA' if i==0 else 'F8FAFD');c._tc.get_or_add_tcPr().append(sh)
        target.add_paragraph().paragraph_format.space_after=Pt(2);return
    if tag=='figure':
        p=target.add_paragraph();inline(p,e);return
    if tag=='pre':
        p=target.add_paragraph(e.text_content());
        for r in p.runs:r.font.size=Pt(8.5)
        return
    for child in e:block(child,target)

for c in chapters:
    doc.add_page_break()
    for e in c['root']:block(e)
doc.save(ROOT/(NAME+'.docx'))

reader=PdfReader(ROOT/(NAME+'.pdf'))
text='\n'.join(p.extract_text() for p in reader.pages)
normalize=lambda s:re.sub(r'\s+','',unicodedata.normalize('NFKC',s))
report={'pdfPages':len(reader.pages),'pdfBytes':(ROOT/(NAME+'.pdf')).stat().st_size,'docxBytes':(ROOT/(NAME+'.docx')).stat().st_size,'chapters':len(chapters),'emptyPages':[i+1 for i,p in enumerate(reader.pages) if len(p.extract_text().strip())<35], 'chapterTextFound':{c['title']:normalize(c['title']) in normalize(text) for c in chapters}}
(ASSETS/'document-build-checks.json').write_text(json.dumps(report,ensure_ascii=False,indent=2))
print(json.dumps(report,ensure_ascii=False))
